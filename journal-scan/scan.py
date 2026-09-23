#!/usr/bin/env python3
"""Scan JM, JMR and JCP for research articles published since 1 January 2025.

Each stage reads the previous stage's JSON in the data directory, so a run can
stop and resume anywhere:

  fetch     Crossref records for each journal                -> candidates.json
  verify    resolve every DOI via the doi.org handle API     -> verified.json
  enrich    abstracts and open-access status, plus future-
            research passages from open-access full text     -> enriched.json
  worklist  print verified papers that still need annotating
  build     merge annotations/*.json into the spreadsheet

DOIs are never typed by hand: every row is keyed from a Crossref record, and
build refuses any annotation whose DOI is not in the doi.org-verified set.

Hosts: api.crossref.org and doi.org are required; api.openalex.org and
api.semanticscholar.org fill in abstracts and open-access status.
"""
from __future__ import annotations

import argparse
import concurrent.futures
import datetime as dt
import html
import io
import json
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

HERE = Path(__file__).resolve().parent
WINDOW_START = dt.date(2025, 1, 1)
# Online-first papers can wait well over a year for an issue, so fetch from
# further back and let in_window() decide.
LOOKBACK = "2023-07-01"
USER_AGENT = "ars-journal-scan/1.0 (academic literature scan)"

JOURNALS = [
    {"key": "JM", "name": "Journal of Marketing", "issn": "0022-2429", "hosts": ("sagepub.com",)},
    {"key": "JMR", "name": "Journal of Marketing Research", "issn": "0022-2437", "hosts": ("sagepub.com",)},
    {"key": "JCP", "name": "Journal of Consumer Psychology", "issn": "1057-7408", "hosts": ("wiley.com",)},
]
JOURNAL_ORDER = {j["key"]: i for i, j in enumerate(JOURNALS)}

HEADERS = ["Title", "Authors", "Journal", "Year/Issue", "DOI link", "Main finding",
           "Method", "Topic tag", "Future research directions", "Interested?"]
INCLUDED_TYPES = {"research", "review"}
KNOWN_TYPES = INCLUDED_TYPES | {"commentary", "rejoinder", "editorial", "other"}
ANNOTATION_FIELDS = ("type", "finding", "method", "tag", "future")

# Titles that are never research articles. Matched against the lowercased title.
EXCLUDE_RULES = [
    (r"^(erratum|corrigend(um|a)|correction\b|retraction|retracted|expression of concern|withdrawn|withdrawal)",
     "correction/retraction notice"),
    (r"\b(editorial board|issue information|masthead|table of contents|front matter|back matter|cover image)\b",
     "front/back matter"),
    (r"\b(call for (papers|proposals)|reviewers? (acknowledg|thank)|thank you to (our )?reviewers|ad hoc reviewers)",
     "journal notice"),
    (r"\b(in memoriam|obituary|book review)\b", "non-research item"),
    (r"^(editorial\b|from the editors?\b|editor[’']?s? (note|introduction|letter|comment)|introduction to the special issue)",
     "editorial"),
]
# Titles that may be commentary rather than research; kept, but shown in the
# worklist so the item type is decided by reading, not by regex.
FLAG_RE = re.compile(r"\b(commentar(y|ies)|rejoinder|reply to|response to|research dialogue|comment on)\b", re.I)
FUTURE_RE = re.compile(r"(future (research|studies|work|inquiry|investigations?)|further research|research agenda"
                       r"|avenues? for (future )?research|limitations and future)", re.I)
INLINE_TAG = re.compile(r"</?(?:jats:|mml:)?(?:i|b|em|strong|sup|sub|scp|sc|italic|bold|span|u|inline-formula"
                        r"|math|mi|mo|mn|mrow)\b[^>]*>", re.I)


# --------------------------------------------------------------------------- http

class NotFound(Exception):
    pass


def get_json(url, *, data=None, tries=4, timeout=60):
    headers = {"User-Agent": USER_AGENT, "Accept": "application/json"}
    body = None
    if data is not None:
        body = json.dumps(data).encode()
        headers["Content-Type"] = "application/json"
    for attempt in range(tries):
        req = urllib.request.Request(url, data=body, headers=headers)
        try:
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                return json.load(resp)
        except urllib.error.HTTPError as e:
            if e.code == 404:
                raise NotFound(url) from None
            if e.code not in (429, 500, 502, 503, 504) or attempt == tries - 1:
                raise
        except (urllib.error.URLError, TimeoutError):
            if attempt == tries - 1:
                raise
        time.sleep(2 ** (attempt + 1))


def get_bytes(url, timeout=60):
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return resp.read(), resp.headers.get_content_type()


# --------------------------------------------------------------------------- parsing

def clean_text(s):
    s = re.sub(r"<jats:title>.*?</jats:title>", " ", s or "", flags=re.S | re.I)
    s = INLINE_TAG.sub("", s)
    s = re.sub(r"<[^>]+>", " ", s)
    s = re.sub(r"\s+", " ", html.unescape(s)).strip()
    return re.sub(r"\s+([,.;:!?)])", r"\1", s)


def cr_date(obj):
    """Crossref {'date-parts': [[y, m, d]]} -> ISO string at the deposited precision."""
    parts = ((obj or {}).get("date-parts") or [[None]])[0]
    if not parts or parts[0] is None:
        return None
    y = int(parts[0])
    m = int(parts[1]) if len(parts) > 1 and parts[1] else None
    d = int(parts[2]) if len(parts) > 2 and parts[2] else None
    return f"{y}" + (f"-{m:02d}" if m else "") + (f"-{d:02d}" if m and d else "")


def iso_to_date(iso):
    y, m, d = ([int(x) for x in iso.split("-")] + [1, 1])[:3]
    return dt.date(y, m, d)


def normalize(rec, journal):
    title = clean_text(" ".join(rec.get("title") or []))
    subtitle = clean_text(" ".join(rec.get("subtitle") or []))
    if subtitle and subtitle.lower() not in title.lower():
        title = f"{title}: {subtitle}"
    authors = []
    for a in rec.get("author") or []:
        name = " ".join(p for p in (a.get("given"), a.get("family")) if p) or a.get("name") or ""
        if name:
            authors.append(clean_text(name))
    volume, issue = rec.get("volume"), rec.get("issue")
    in_issue = bool(volume and issue)
    issue_obj = (rec.get("journal-issue") or {}).get("published-print") or rec.get("published-print")
    online = cr_date(rec.get("published-online"))
    c = {
        "doi": rec["DOI"].lower(),
        "journal_key": journal["key"],
        "journal": journal["name"],
        "title": title,
        "authors": authors,
        "volume": volume,
        "issue": issue,
        "page": rec.get("page") or rec.get("article-number"),
        "online_date": online or (None if in_issue else cr_date(rec.get("issued"))),
        "issue_date": (cr_date(issue_obj) or cr_date(rec.get("issued"))) if in_issue else None,
        "in_issue": in_issue,
        "crossref_abstract": clean_text(rec.get("abstract") or "") or None,
        "is_update_notice": bool(rec.get("update-to")),
    }
    c["year_issue"] = year_issue(c)
    c["exclude_reason"] = classify(c)
    c["flags"] = sorted({m.group(0).lower() for m in FLAG_RE.finditer(title)})
    return c


def year_issue(c):
    if c["in_issue"]:
        year = (c["issue_date"] or c["online_date"] or "")[:4]
        return f"{year}, Vol. {c['volume']}, Issue {c['issue']}".lstrip(", ")
    return "Online first"


def classify(c):
    """Reason a record is not a research article, or None."""
    if c["is_update_notice"]:
        return "correction/retraction notice"
    title = c["title"].lower()
    for pattern, reason in EXCLUDE_RULES:
        if re.search(pattern, title):
            return reason
    if not c["authors"]:
        return "no authors (front/back matter)"
    return None


def in_window(c, today):
    """In-issue papers count by issue date (so a 2024 online-first paper printed
    in a 2025 issue is in); online-first papers count by online date."""
    if c["in_issue"] and c["issue_date"]:
        first_public = min(iso_to_date(d) for d in (c["online_date"], c["issue_date"]) if d)
        return iso_to_date(c["issue_date"]) >= WINDOW_START and first_public <= today
    if not c["online_date"]:
        return False
    return WINDOW_START <= iso_to_date(c["online_date"]) <= today


def sort_key(c):
    def num(v):
        m = re.match(r"\d+", str(v or ""))
        return int(m.group()) if m else 0
    return (JOURNAL_ORDER[c["journal_key"]], 0 if c["in_issue"] else 1, num(c["volume"]), num(c["issue"]),
            num(c["page"]), c["online_date"] or "", c["doi"])


# --------------------------------------------------------------------------- io

def read_json(path):
    return json.loads(Path(path).read_text())


def write_json(path, obj):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(obj, indent=1, ensure_ascii=False) + "\n")


def load_annotations(data_dir):
    merged = {}
    for f in sorted((Path(data_dir) / "annotations").glob("*.json")):
        for doi, a in read_json(f).items():
            doi = doi.lower()
            if doi in merged:
                sys.exit(f"{f.name}: {doi} is annotated twice")
            merged[doi] = a
    return merged


# --------------------------------------------------------------------------- fetch

def fetch_crossref(journal):
    base = f"https://api.crossref.org/journals/{journal['issn']}/works"
    items, cursor = [], "*"
    while True:
        q = urllib.parse.urlencode({"filter": f"from-pub-date:{LOOKBACK}", "rows": 500, "cursor": cursor})
        msg = get_json(f"{base}?{q}")["message"]
        items += msg["items"]
        if not msg["items"] or len(items) >= msg["total-results"]:
            return items
        cursor = msg["next-cursor"]


def cmd_fetch(args):
    today = dt.date.today()
    out, seen = [], set()
    for j in JOURNALS:
        records = fetch_crossref(j)
        kept = excluded = 0
        for rec in records:
            if rec.get("type") != "journal-article" or rec["DOI"].lower() in seen:
                continue
            seen.add(rec["DOI"].lower())
            c = normalize(rec, j)
            if not in_window(c, today):
                continue
            out.append(c)
            kept += 1
            excluded += bool(c["exclude_reason"])
        print(f"{j['key']}: {len(records)} Crossref records since {LOOKBACK}, {kept} in window "
              f"({excluded} non-research auto-excluded)")
    out.sort(key=sort_key)
    write_json(Path(args.data) / "candidates.json", {
        "fetched_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "window": [WINDOW_START.isoformat(), today.isoformat()],
        "items": out,
    })


# --------------------------------------------------------------------------- verify

def check_doi(doi):
    url = "https://doi.org/api/handles/" + urllib.parse.quote(doi, safe="/") + "?type=URL"
    try:
        data = get_json(url)
    except NotFound:
        return {"verified": False, "verify_reason": "doi.org: DOI not registered"}
    except Exception as e:  # noqa: BLE001 — any transport failure means "could not check"
        return {"verified": None, "verify_reason": f"doi.org unreachable ({type(e).__name__}: {e})"}
    if data.get("responseCode") != 1:
        return {"verified": False, "verify_reason": f"doi.org responseCode {data.get('responseCode')}"}
    target = next((v["data"]["value"] for v in data.get("values", []) if v.get("type") == "URL"), None)
    if not target:
        return {"verified": False, "verify_reason": "doi.org: no URL registered"}
    return {"verified": True, "resolves_to": target}


def cmd_verify(args):
    data = read_json(Path(args.data) / "candidates.json")
    todo = [c for c in data["items"] if not c["exclude_reason"]]
    with concurrent.futures.ThreadPoolExecutor(max_workers=6) as pool:
        for c, result in zip(todo, pool.map(check_doi, (c["doi"] for c in todo))):
            c.update(result)
            host = urllib.parse.urlparse(result.get("resolves_to") or "").hostname or ""
            journal = next(j for j in JOURNALS if j["key"] == c["journal_key"])
            if result["verified"] and not host.endswith(journal["hosts"]):
                c["verify_note"] = f"resolves outside the publisher's domain ({host})"
    data["verified_at"] = dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")
    write_json(Path(args.data) / "verified.json", data)
    ok = sum(c.get("verified") is True for c in todo)
    failed = [c for c in todo if c.get("verified") is not True]
    print(f"verified {ok}/{len(todo)} DOIs")
    for c in failed:
        print(f"  NOT VERIFIED {c['doi']}: {c['verify_reason']}")


# --------------------------------------------------------------------------- enrich

def from_inverted(index):
    words = sorted((pos, word) for word, positions in index.items() for pos in positions)
    return " ".join(w for _, w in words)


def openalex_lookup(dois):
    out = {}
    for i in range(0, len(dois), 40):
        q = urllib.parse.urlencode({"filter": "doi:" + "|".join(dois[i:i + 40]), "per-page": 50,
                                    "select": "doi,abstract_inverted_index,open_access"})
        try:
            data = get_json(f"https://api.openalex.org/works?{q}")
        except Exception as e:  # noqa: BLE001
            print(f"  OpenAlex batch {i // 40 + 1} failed: {e}", file=sys.stderr)
            continue
        for w in data.get("results", []):
            doi = (w.get("doi") or "").lower().removeprefix("https://doi.org/")
            oa = w.get("open_access") or {}
            idx = w.get("abstract_inverted_index")
            out[doi] = {"abstract": clean_text(from_inverted(idx)) if idx else None, "is_oa": oa.get("is_oa"),
                        "oa_status": oa.get("oa_status"), "oa_url": oa.get("oa_url")}
        time.sleep(0.2)
    return out


def s2_lookup(dois):
    out = {}
    for i in range(0, len(dois), 400):
        chunk = dois[i:i + 400]
        try:
            papers = get_json("https://api.semanticscholar.org/graph/v1/paper/batch?fields=abstract,openAccessPdf",
                              data={"ids": [f"DOI:{d}" for d in chunk]})
        except Exception as e:  # noqa: BLE001
            print(f"  Semantic Scholar batch failed: {e}", file=sys.stderr)
            continue
        for doi, p in zip(chunk, papers):
            if p:
                out[doi] = {"abstract": clean_text(p.get("abstract") or "") or None,
                            "oa_url": (p.get("openAccessPdf") or {}).get("url")}
    return out


def future_passages(url):
    try:
        body, ctype = get_bytes(url)
    except Exception as e:  # noqa: BLE001
        return [], f"full text not retrievable ({type(e).__name__})"
    if ctype == "application/pdf" or body[:4] == b"%PDF":
        try:
            from pypdf import PdfReader
        except ImportError:
            return [], "pypdf not installed"
        text = " ".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(body)).pages)
    else:
        text = clean_text(re.sub(r"(?is)<(script|style)\b.*?</\1>", " ", body.decode("utf-8", "replace")))
    text = re.sub(r"\s+", " ", text)
    refs = max(text.rfind("References"), text.rfind("REFERENCES"))
    if refs > len(text) // 2:
        text = text[:refs]  # a reference list is full of titles that match FUTURE_RE
    sentences = re.split(r"(?<=[.!?])\s+(?=[A-Z])", text)
    hits = [" ".join(sentences[i:i + 3])[:700] for i, s in enumerate(sentences) if FUTURE_RE.search(s)]
    return hits[-4:], ("ok" if hits else "no future-research passage found in full text")


def cmd_enrich(args):
    data = read_json(Path(args.data) / "verified.json")
    target = [c for c in data["items"] if c.get("verified") is True and not c["exclude_reason"]]
    oa = openalex_lookup([c["doi"] for c in target])
    missing = [c["doi"] for c in target if not (c["crossref_abstract"] or (oa.get(c["doi"]) or {}).get("abstract"))]
    s2 = s2_lookup(missing) if missing else {}
    for c in target:
        o, s = oa.get(c["doi"]) or {}, s2.get(c["doi"]) or {}
        for source, text in (("Crossref (publisher deposit)", c["crossref_abstract"]),
                             ("OpenAlex", o.get("abstract")), ("Semantic Scholar", s.get("abstract"))):
            if text:
                c["abstract"], c["abstract_source"] = text, source
                break
        else:
            c["abstract"], c["abstract_source"] = None, None
        c["is_oa"] = o["is_oa"] if "is_oa" in o else (True if s.get("oa_url") else None)
        c["oa_status"] = o.get("oa_status")
        c["oa_url"] = o.get("oa_url") or s.get("oa_url")
        if args.fulltext and c["oa_url"]:
            c["future_passages"], c["fulltext_note"] = future_passages(c["oa_url"])
    data["enriched_at"] = dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")
    write_json(Path(args.data) / "enriched.json", data)
    print(f"{len(target)} papers: {sum(bool(c['abstract']) for c in target)} with abstracts, "
          f"{sum(c['is_oa'] is True for c in target)} open access, "
          f"{sum(c['is_oa'] is None for c in target)} access unknown")


# --------------------------------------------------------------------------- worklist

def cmd_worklist(args):
    items = [c for c in read_json(Path(args.data) / "enriched.json")["items"]
             if c.get("verified") is True and not c["exclude_reason"]]
    done = load_annotations(args.data)
    todo = [c for c in items if c["doi"] not in done and (not args.journal or c["journal_key"] == args.journal)]
    for c in todo[:args.batch]:
        print(f"### {c['doi']} | {c['journal_key']} | {c['year_issue']}")
        print(f"TITLE: {c['title']}")
        if c["flags"]:
            print(f"FLAGS: {', '.join(c['flags'])}")
        print(f"ABSTRACT [{c.get('abstract_source') or 'none'}]: {c.get('abstract') or '(not retrievable)'}")
        for p in c.get("future_passages") or []:
            print(f"FUTURE: {p}")
        print()
    print(f"-- {len(todo)} to annotate, {len(done)} done, {len(items)} verified papers", file=sys.stderr)


# --------------------------------------------------------------------------- build

def validate_annotations(ann):
    problems = []
    for doi, a in ann.items():
        missing = [f for f in ANNOTATION_FIELDS if not str(a.get(f) or "").strip()]
        if missing:
            problems.append(f"{doi}: missing {', '.join(missing)}")
            continue
        if a["type"] not in KNOWN_TYPES:
            problems.append(f"{doi}: unknown type {a['type']!r}")
        if a["type"] in INCLUDED_TYPES and len(a["tag"].split()) > 2:
            problems.append(f"{doi}: topic tag {a['tag']!r} is longer than two words")
    return problems


def access_label(c):
    if c.get("is_oa") is True:
        return f"Open access ({c['oa_status']})" if c.get("oa_status") else "Open access"
    if c.get("is_oa") is False:
        return "Subscription only"
    return "Unknown"


def cmd_build(args):
    data = read_json(Path(args.data) / "enriched.json")
    items = data["items"]
    ann = load_annotations(args.data)
    verified = {c["doi"] for c in items if c.get("verified") is True and not c["exclude_reason"]}

    unknown = sorted(set(ann) - verified)
    if unknown:
        sys.exit("annotations name DOIs that are not verified candidates (typo or invented?):\n  "
                 + "\n  ".join(unknown))
    missing = sorted(verified - set(ann))
    if missing and not args.allow_missing:
        sys.exit(f"{len(missing)} verified papers have no annotation yet; run `worklist`")
    problems = validate_annotations(ann)
    if problems:
        sys.exit("annotation problems:\n  " + "\n  ".join(problems))

    papers, excluded, limited = [], [], []
    for c in items:
        if c["exclude_reason"]:
            excluded.append((c, c["exclude_reason"]))
            continue
        if c.get("verified") is not True:
            limited.append((c, "No", c.get("verify_reason") or "DOI not verified"))
            continue
        a = ann.get(c["doi"]) or {"type": "research", "finding": "PENDING", "method": "PENDING",
                                   "tag": "PENDING", "future": "PENDING"}
        if a["type"] not in INCLUDED_TYPES:
            excluded.append((c, a["type"]))
            continue
        papers.append((c, a))
        notes = []
        if not c.get("abstract"):
            notes.append("no abstract retrievable, so the main finding is not available")
        if c.get("is_oa") is False:
            notes.append("full text is subscription-only")
        elif c.get("is_oa") is None:
            notes.append("open-access status unknown")
        if notes:
            limited.append((c, "Yes", "; ".join(notes)))

    write_workbook(Path(args.out), data, papers, excluded, limited)
    tags = sorted({a["tag"] for _, a in papers})
    print(f"wrote {args.out}: {len(papers)} papers, {len(tags)} topic tags, "
          f"{len(limited)} limited-access entries, {len(excluded)} excluded items")


def write_workbook(out, data, papers, excluded, limited):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    base = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    link_font = Font(name="Arial", size=10, color="0563C1", underline="single")
    fill_in = PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")

    def table(ws, headers, rows, widths, link_col=None):
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill, cell.alignment = head_font, head_fill, wrap
        for row in rows:
            ws.append(row)
        for r in ws.iter_rows(min_row=2):
            for cell in r:
                cell.font, cell.alignment = base, wrap
            if link_col is not None and r[link_col].value:
                r[link_col].hyperlink = r[link_col].value
                r[link_col].font = link_font
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions

    doi_url = "https://doi.org/{}".format
    wb = Workbook()

    ws = wb.active
    ws.title = "Papers"
    table(ws, HEADERS, [[c["title"], ", ".join(c["authors"]), c["journal"], c["year_issue"], doi_url(c["doi"]),
                         a["finding"], a["method"], a["tag"], a["future"], None] for c, a in papers],
          [48, 32, 24, 22, 36, 60, 28, 16, 46, 12], link_col=4)
    for (cell,) in ws.iter_rows(min_row=2, min_col=10, max_col=10):
        cell.fill = fill_in
    if papers:
        dv = DataValidation(type="list", formula1='"Yes,Maybe,No"', allow_blank=True, showErrorMessage=False)
        dv.add(f"J2:J{len(papers) + 1}")
        ws.add_data_validation(dv)

    ws = wb.create_sheet("Topic counts")
    tags = sorted({a["tag"] for _, a in papers}, key=str.lower)
    names = [j["name"] for j in JOURNALS]
    ws.append(["Topic tag", *names, "Total"])
    for i, tag in enumerate(tags, start=2):
        ws.append([tag, *[f'=COUNTIFS(Papers!$H:$H,$A{i},Papers!$C:$C,{col}$1)' for col in "BCD"],
                   f"=COUNTIF(Papers!$H:$H,$A{i})"])
    last = len(tags) + 1
    ws.append(["Total", *[f"=SUM({col}2:{col}{last})" for col in "BCDE"]])
    ws.append(["Rows in Papers", *[f'=COUNTIF(Papers!$C:$C,{col}$1)' for col in "BCD"], "=COUNTA(Papers!$A:$A)-1"])
    ws.append(["Check", *[f'=IF({col}{last + 1}={col}{last + 2},"OK","MISMATCH")' for col in "BCDE"]])
    for r in ws.iter_rows():
        for cell in r:
            cell.font = base
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment = head_font, head_fill, wrap
    for r in (last + 1, last + 2, last + 3):
        for cell in ws[r]:
            cell.font = bold
    for col, w in zip("ABCDE", (24, 20, 20, 20, 10)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Limited access")
    table(ws, ["Title", "Journal", "Year/Issue", "DOI link", "In Papers sheet?", "Access", "What is limited"],
          [[c["title"], c["journal"], c["year_issue"], doi_url(c["doi"]), in_papers, access_label(c), why]
           for c, in_papers, why in limited],
          [48, 24, 22, 36, 12, 20, 50], link_col=3)

    ws = wb.create_sheet("Excluded")
    table(ws, ["Title", "Journal", "Year/Issue", "DOI link", "Why excluded"],
          [[c["title"], c["journal"], c["year_issue"], doi_url(c["doi"]), why] for c, why in excluded],
          [60, 24, 22, 36, 30], link_col=3)

    ws = wb.create_sheet("About")
    start, end = data["window"]
    about = [
        ("Scope", f"Research articles in Journal of Marketing, Journal of Marketing Research and Journal of "
                  f"Consumer Psychology published {start} to {end}, including online-first / Early View papers."),
        ("Papers listed", "=COUNTA(Papers!A:A)-1"),
        *[(f"  {j['name']}", f'=COUNTIF(Papers!C:C,"{j["name"]}")') for j in JOURNALS],
        ("Metadata", f"Crossref REST API (publisher-deposited records), retrieved {data['fetched_at']}. "
                     f"ISSNs: " + "; ".join(f"{j['key']} {j['issn']}" for j in JOURNALS) + "."),
        ("Date rule", "A paper in a numbered issue counts if the issue is dated 2025 or later (so a 2024 "
                      "online-first paper printed in a 2025 issue is in). A paper not yet in an issue counts "
                      "if it went online on or after 1 January 2025; its Year/Issue reads \"Online first\"."),
        ("DOI verification", f"Every DOI in Papers resolved through the doi.org handle API "
                             f"(responseCode 1) on {data.get('verified_at', 'n/a')}. Candidates whose DOI could "
                             f"not be verified are listed in Limited access, never in Papers. No DOI was typed "
                             f"by hand: each comes from its Crossref record."),
        ("Excluded items", "Editorials, corrections, retractions, front/back matter and journal notices were "
                           "removed by title rules; commentaries and rejoinders were removed after reading. "
                           "All are listed in the Excluded sheet."),
        ("Summaries", "Main finding, Method, Topic tag and Future research directions are written in plain "
                      "English from each paper's abstract (and open-access full text where available). They "
                      "paraphrase and do not quote. Papers are not ranked or recommended; rows are ordered by "
                      "journal, then issue, then online date."),
        ("Future research", "\"not available\" means neither the abstract nor an openly accessible full text "
                            "states future research directions. Most full texts in these journals are "
                            "subscription-only; see Limited access."),
        ("Interested?", "Left blank for you. Each cell offers Yes / Maybe / No from a dropdown, but accepts "
                        "any text."),
        ("Topic counts", "Counts are live formulas over the Papers sheet, so they update if you edit a tag."),
        ("AI assistance", "Compiled with AI assistance (Claude, using the academic-research-skills deep-research "
                          "skill). Check key details on the DOI landing page before citing a paper."),
    ]
    for row in about:
        ws.append(list(row))
    for r in ws.iter_rows():
        r[0].font, r[1].font = bold, base
        r[0].alignment = r[1].alignment = wrap
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


# --------------------------------------------------------------------------- cli

def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--data", default=str(HERE / "data"), help="data directory (default: %(default)s)")
    sub = p.add_subparsers(dest="cmd", required=True)
    sub.add_parser("fetch").set_defaults(fn=cmd_fetch)
    sub.add_parser("verify").set_defaults(fn=cmd_verify)
    e = sub.add_parser("enrich")
    e.add_argument("--fulltext", action="store_true", help="mine open-access full text for future-research passages")
    e.set_defaults(fn=cmd_enrich)
    w = sub.add_parser("worklist")
    w.add_argument("--batch", type=int, default=25)
    w.add_argument("--journal", choices=[j["key"] for j in JOURNALS])
    w.set_defaults(fn=cmd_worklist)
    b = sub.add_parser("build")
    b.add_argument("--out", default=str(HERE / "journal_scan.xlsx"))
    b.add_argument("--allow-missing", action="store_true", help="write PENDING for unannotated papers")
    b.set_defaults(fn=cmd_build)
    args = p.parse_args(argv)
    args.fn(args)


if __name__ == "__main__":
    main()
