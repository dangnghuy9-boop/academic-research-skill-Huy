"""Step 4: build the workbook from data/papers.json + data/summaries.json.

Sheets: Papers (user's exact 10 columns) | Topic counts (COUNTIF formulas) | Access notes | About
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from common import JOURNALS, START, TODAY, DATA, load

HEADERS = ["Title", "Authors", "Journal", 'Year/Issue (or "Online first")', "DOI link",
           "Main finding", "Method", "Topic tag", "Future research directions", "Interested?"]
WIDTHS = [48, 30, 22, 22, 34, 60, 30, 16, 50, 12]
F = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # the one column the reader fills in
THIN = Side(style="thin", color="BFBFBF")
JORDER = {k: i for i, k in enumerate(JOURNALS)}


def sort_key(p):
    # Neutral ordering (journal, then publication chronology) — not a ranking.
    in_issue = 0 if p["in_issue"] else 1
    d = p["print_date"] or p["issued"] or p["online_date"] or ""
    return (JORDER[p["journal_key"]], in_issue, p["volume"] or "", d, p["online_date"] or "", p["title"].lower())


def header(ws, cols, widths):
    for c, (h, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=F, bold=True, color="FFFFFF")
        cell.fill = HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def main(out_path):
    papers = sorted([p for p in load(DATA / "papers.json") if p["include"]], key=sort_key)
    summ = load(DATA / "summaries.json", {})
    missing = [p["doi"] for p in papers if p["doi"] not in summ]
    if missing:
        sys.exit(f"{len(missing)} included papers lack summaries, e.g. {missing[:5]}")

    wb = Workbook()
    ws = wb.active
    ws.title = "Papers"
    header(ws, HEADERS, WIDTHS)
    body = Font(name=F, size=10)
    for r, p in enumerate(papers, 2):
        s = summ[p["doi"]]
        url = f"https://doi.org/{p['doi']}"
        vals = [p["title"], p["authors"], p["journal"], p["year_issue"], url,
                s["finding"], s["method"], s["tag"], s["future"], None]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = body
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN)
        link = ws.cell(row=r, column=5)
        link.hyperlink = url
        link.font = Font(name=F, size=10, color="0563C1", underline="single")
        ws.cell(row=r, column=10).fill = INPUT_FILL
    last = len(papers) + 1
    ws.auto_filter.ref = f"A1:J{last}"
    dv = DataValidation(type="list", formula1='"Yes,Maybe,No"', allow_blank=True)
    dv.prompt, dv.promptTitle = "Optional: pick Yes / Maybe / No, or leave blank.", "Interested?"
    ws.add_data_validation(dv)
    dv.add(f"J2:J{last}")

    # ---- Topic counts: live formulas against the Papers sheet ----
    tc = wb.create_sheet("Topic counts")
    jnames = [JOURNALS[k]["name"] for k in JOURNALS]
    header(tc, ["Topic tag", "All three journals"] + jnames, [26, 18, 22, 26, 28])
    tags = sorted({summ[p["doi"]]["tag"] for p in papers}, key=str.lower)
    rng_tag, rng_j = f"Papers!$H$2:$H${last}", f"Papers!$C$2:$C${last}"
    for r, t in enumerate(tags, 2):
        tc.cell(row=r, column=1, value=t).font = body
        tc.cell(row=r, column=2, value=f"=COUNTIF({rng_tag},$A{r})").font = body
        for c, jn in enumerate(jnames, 3):
            tc.cell(row=r, column=c, value=f'=COUNTIFS({rng_tag},$A{r},{rng_j},"{jn}")').font = body
    tr = len(tags) + 2
    tc.cell(row=tr, column=1, value="Total").font = Font(name=F, bold=True)
    for c in range(2, 6):
        col = get_column_letter(c)
        tc.cell(row=tr, column=c, value=f"=SUM({col}2:{col}{tr-1})").font = Font(name=F, bold=True)
        tc.cell(row=tr, column=c).border = Border(top=Side(style="thin"))
    tc.cell(row=tr + 1, column=1, value="Check: rows on Papers sheet").font = Font(name=F, italic=True, size=9)
    tc.cell(row=tr + 1, column=2, value=f"=COUNTA(Papers!$A$2:$A${last})").font = Font(name=F, italic=True, size=9)
    tc.cell(row=tr + 2, column=1, value="Tags are sorted alphabetically; counts update if you retag a paper on the Papers sheet.").font = Font(name=F, italic=True, size=9)

    # ---- Access notes ----
    an = wb.create_sheet("Access notes")
    header(an, ["Title", "Journal", "DOI link", "Access", "What was limited", "Effect on this sheet"],
           [48, 22, 34, 16, 44, 50])
    notes = [(p, summ[p["doi"]]) for p in papers if summ[p["doi"]].get("access_note")]
    for r, (p, s) in enumerate(notes, 2):
        url = f"https://doi.org/{p['doi']}"
        access = "Open access" if p.get("is_oa") else "Subscription"
        for c, v in enumerate([p["title"], p["journal"], url, access, s["access_note"], s.get("access_effect", "")], 1):
            cell = an.cell(row=r, column=c, value=v)
            cell.font = body
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        an.cell(row=r, column=3).hyperlink = url
        an.cell(row=r, column=3).font = Font(name=F, size=10, color="0563C1", underline="single")
    if not notes:
        an.cell(row=2, column=1, value="No access limitations recorded.").font = body

    # ---- About ----
    ab = wb.create_sheet("About")
    ab.column_dimensions["A"].width = 30
    ab.column_dimensions["B"].width = 110
    about = load(DATA / "about.json", [])
    for r, (k, v) in enumerate(about, 1):
        ab.cell(row=r, column=1, value=k).font = Font(name=F, bold=True, size=10)
        c = ab.cell(row=r, column=2, value=v)
        c.font = Font(name=F, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(out_path)
    print(f"saved {out_path}: {len(papers)} papers, {len(tags)} tags, {len(notes)} access notes")


if __name__ == "__main__":
    main(sys.argv[1])
